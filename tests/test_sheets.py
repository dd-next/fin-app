"""Google Sheets sync tests — no network, gspread is faked throughout."""

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import gspread
import pytest
from openpyxl import load_workbook

from app import export, sheets

D = Decimal


# ---- fakes -------------------------------------------------------------------


class FakeWorksheet:
    def __init__(self):
        self.rows = None
        self.cleared = False

    def clear(self):
        self.cleared = True

    def update(self, rows):
        self.rows = rows


class FakeSpreadsheet:
    def __init__(self):
        self.worksheets = {}

    def worksheet(self, title):
        if title not in self.worksheets:
            raise gspread.WorksheetNotFound(title)
        return self.worksheets[title]

    def add_worksheet(self, title, rows, cols):
        self.worksheets[title] = FakeWorksheet()
        return self.worksheets[title]


@pytest.fixture
def fake_sheet(monkeypatch):
    """Enable Sheets and capture writes in a FakeSpreadsheet."""
    monkeypatch.setenv("SHEETS_ENABLED", "true")
    monkeypatch.setenv("GOOGLE_SHEET_ID", "fake-sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "/nonexistent.json")
    fake = FakeSpreadsheet()
    monkeypatch.setattr(sheets, "_open_spreadsheet", lambda: fake)
    return fake


def sample_period(days=10):
    start = date.today()
    return {
        "total_amount": "1000",
        "start_date": start.isoformat(),
        "end_date": (start + timedelta(days=days - 1)).isoformat(),
    }


# ---- shared row-building: export and sync agree -------------------------------


def test_workbook_matches_shared_row_functions():
    start, end = date(2026, 7, 1), date(2026, 7, 10)
    rows = [
        export.OperationRow(datetime(2026, 7, 1, 9), D("250"), "groceries"),
        export.OperationRow(datetime(2026, 7, 2, 9), D("100"), None, "income"),
    ]
    wb = load_workbook(
        BytesIO(export.build_workbook(D("1000"), start, end, rows, today=start))
    )

    period_sheet = [list(r) for r in wb["Period"].iter_rows(values_only=True)]
    expected_period = export.period_rows(D("1000"), start, end, rows, today=start)
    assert [r[0] for r in period_sheet] == [r[0] for r in expected_period]
    for got, want in zip(period_sheet, expected_period):
        # numbers compare numerically, strings (dates) literally
        if isinstance(want[1], (D, int)):
            assert D(str(got[1])) == D(str(want[1]))
        else:
            assert got[1] == want[1]

    expense_sheet = [list(r) for r in wb["Expenses"].iter_rows(values_only=True)]
    assert expense_sheet[0] == export.EXPENSE_HEADERS
    expected_rows = export.expense_rows(D("1000"), start, end, rows)
    assert len(expense_sheet[1:]) == len(expected_rows)
    for got, want in zip(expense_sheet[1:], expected_rows):
        assert got[0] == want[0]  # date
        assert D(str(got[1])) == want[1]  # amount
        assert got[2] == want[2]  # type (expense | income)
        assert (got[3] or "") == want[3]  # comment
        assert D(str(got[4])) == want[4]  # running balance
        assert D(str(got[5])) == want[5]  # per-day allowance


# ---- mutations trigger the sync ------------------------------------------------


async def test_mutation_writes_expected_rows(client, fake_sheet):
    resp = await client.post("/period", json=sample_period())
    assert resp.status_code == 200

    resp = await client.post("/operations", json={"amount": "250", "comment": "food"})
    assert resp.status_code == 200

    period_ws = fake_sheet.worksheets["Period"]
    assert period_ws.cleared
    kv = {row[0]: row[1] for row in period_ws.rows}
    assert kv["total_amount"] == 1000.0
    assert kv["days_total"] == 10
    assert kv["spent_total"] == 250.0
    assert kv["remaining"] == 750.0

    expenses_ws = fake_sheet.worksheets["Expenses"]
    assert expenses_ws.rows[0] == export.EXPENSE_HEADERS
    body = expenses_ws.rows[1:]
    assert len(body) == 1
    assert body[0][1] == 250.0
    assert body[0][2] == "expense"
    assert body[0][3] == "food"
    assert body[0][4] == 750.0
    assert body[0][5] == -150.0  # today's 100 budget - 250 spent, 1:1


async def test_delete_resyncs(client, fake_sheet):
    await client.post("/period", json=sample_period())
    resp = await client.post("/operations", json={"amount": "100"})
    operation_id = resp.json()["operation"]["id"]

    resp = await client.delete(f"/operations/{operation_id}")
    assert resp.status_code == 200

    # sheet mirrors the DB again: no expense rows, full amount remaining
    assert fake_sheet.worksheets["Expenses"].rows == [export.EXPENSE_HEADERS]
    kv = {row[0]: row[1] for row in fake_sheet.worksheets["Period"].rows}
    assert kv["remaining"] == 1000.0


# ---- best-effort: failures never break the API ---------------------------------


async def test_gspread_error_swallowed(client, monkeypatch):
    monkeypatch.setenv("SHEETS_ENABLED", "true")
    monkeypatch.setenv("GOOGLE_SHEET_ID", "fake-sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "/nonexistent.json")

    def explode():
        raise RuntimeError("Google is down")

    monkeypatch.setattr(sheets, "_open_spreadsheet", explode)

    resp = await client.post("/period", json=sample_period())
    assert resp.status_code == 200
    resp = await client.post("/operations", json={"amount": "10"})
    assert resp.status_code == 200  # sync failed silently, API unaffected


async def test_disabled_makes_zero_google_calls(client, monkeypatch):
    # SHEETS_ENABLED defaults to false; make any Google touch blow the test.
    monkeypatch.delenv("SHEETS_ENABLED", raising=False)

    def forbidden():
        raise AssertionError("Google API must not be called when disabled")

    monkeypatch.setattr(sheets, "_open_spreadsheet", forbidden)

    await client.post("/period", json=sample_period())
    resp = await client.post("/operations", json={"amount": "10"})
    assert resp.status_code == 200


# ---- manual sync endpoint -------------------------------------------------------


async def test_manual_sync_ok(client, fake_sheet):
    await client.post("/period", json=sample_period())
    resp = await client.post("/sheets/sync")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}
    assert "Period" in fake_sheet.worksheets


async def test_manual_sync_failed(client, monkeypatch):
    monkeypatch.setenv("SHEETS_ENABLED", "true")
    monkeypatch.setenv("GOOGLE_SHEET_ID", "fake-sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "/nonexistent.json")

    def explode():
        raise RuntimeError("Google is down")

    monkeypatch.setattr(sheets, "_open_spreadsheet", explode)

    await client.post("/period", json=sample_period())
    resp = await client.post("/sheets/sync")
    assert resp.status_code == 200
    assert resp.json() == {"status": "failed"}


async def test_manual_sync_disabled(client, monkeypatch):
    monkeypatch.delenv("SHEETS_ENABLED", raising=False)
    resp = await client.post("/sheets/sync")
    assert resp.json() == {"status": "disabled"}
