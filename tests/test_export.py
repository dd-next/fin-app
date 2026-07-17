"""Round-trip test: create data via the API, download /export.xlsx, load the
bytes back with openpyxl, and check both sheets against the API's numbers."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

D = Decimal


async def test_export_roundtrip(client):
    start = date.today()
    end = start + timedelta(days=9)  # 10-day period
    resp = await client.post(
        "/api/v1/workspaces/1/periods",
        json={
            "total_amount": "1000",
            "start_date": start.isoformat(),
            "end_date": end.isoformat(),
        },
    )
    assert resp.status_code == 200
    period_id = resp.json()["period"]["id"]

    for amount, comment in [("250", "groceries"), ("100", None), ("50.50", "coffee")]:
        resp = await client.post("/api/v1/workspaces/1/operations", json={"amount": amount, "comment": comment})
        assert resp.status_code == 200

    api_budget = (await client.get("/api/v1/workspaces/1/budget")).json()
    api_operations = (await client.get("/api/v1/workspaces/1/operations")).json()

    resp = await client.get(f"/api/v1/workspaces/1/periods/{period_id}/export.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.headers["content-disposition"] == 'attachment; filename="finapp-export.xlsx"'

    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Period", "Expenses"]

    # --- Period sheet: key/value pairs match the API ------------------------
    period_kv = {row[0]: row[1] for row in wb["Period"].iter_rows(values_only=True)}
    assert D(str(period_kv["total_amount"])) == D("1000")
    assert period_kv["start_date"] == start.isoformat()
    assert period_kv["end_date"] == end.isoformat()
    assert period_kv["days_total"] == api_budget["days_total"]
    assert D(str(period_kv["spent_total"])) == D(api_budget["spent_total"]) == D("400.50")
    assert D(str(period_kv["remaining"])) == D(api_budget["remaining_money"]) == D("599.50")

    # --- Expenses sheet: header + one row per operation, replayed in order --
    rows = list(wb["Expenses"].iter_rows(values_only=True))
    assert rows[0] == (
        "Date",
        "Amount",
        "Type",
        "Comment",
        "Running balance",
        "Left to spend that day",
        "Category",
        "Pool",
        "Savings goal",
        "Added by",
    )
    body = rows[1:]
    assert len(body) == len(api_operations) == 3

    # chronological order (API lists newest first — export replays oldest first)
    amounts = [D(str(r[1])) for r in body]
    assert amounts == [D("250"), D("100"), D("50.50")]
    assert sum(amounts) == D(api_budget["spent_total"])
    assert [r[2] for r in body] == ["expense", "expense", "expense"]

    # running balance replays correctly and ends at the API's remaining_money
    balances = [D(str(r[4])) for r in body]
    assert balances == [D("750"), D("650"), D("599.50")]
    assert balances[-1] == D(api_budget["remaining_money"])

    # "left to spend that day" at the last point matches the API's headline
    # number (all expenses were added today)
    assert D(str(body[-1][5])) == D(api_budget["per_day_today"])

    # dates are YYYY-MM-DD strings; comments preserved (blank for None —
    # openpyxl reads empty cells back as None)
    assert all(r[0] == start.isoformat() for r in body)
    assert [r[3] or "" for r in body] == ["groceries", "", "coffee"]
    assert [r[6] for r in body] == ["Uncategorized"] * 3


async def test_export_includes_category_pool_and_author(client, monkeypatch):
    monkeypatch.setenv("WEB_AUTH_ENABLED", "true")
    monkeypatch.setenv("BOOTSTRAP_TOKEN", "setup-secret")
    owner = await client.post(
        "/api/v1/auth/bootstrap",
        json={
            "username": "owner",
            "password": "owner-password-123",
            "display_name": "Alex",
        },
        headers={"X-Bootstrap-Token": "setup-secret"},
    )
    assert owner.status_code == 200

    start = date.today()
    period = await client.post(
        "/api/v1/workspaces/1/periods",
        json={
            "total_amount": "1000",
            "start_date": start.isoformat(),
            "end_date": (start + timedelta(days=9)).isoformat(),
        },
    )
    period_id = period.json()["period"]["id"]
    category = (
        await client.post(
            "/api/v1/workspaces/1/categories", json={"name": "Rent"}
        )
    ).json()
    pool = (
        await client.post("/api/v1/workspaces/1/pools", json={"name": "Home"})
    ).json()
    pool_plan = (
        await client.put(
            f"/api/v1/workspaces/1/periods/{period_id}/pool-plans/{pool['id']}",
            json={"allocated_amount": "500"},
        )
    ).json()
    await client.put(
        f"/api/v1/workspaces/1/periods/{period_id}/category-plans/{category['id']}",
        json={"limit_amount": "500", "pool_plan_id": pool_plan["id"]},
    )
    await client.post(
        f"/api/v1/workspaces/1/periods/{period_id}/operations",
        json={"amount": "100", "category_id": category["id"]},
    )

    response = await client.get(
        f"/api/v1/workspaces/1/periods/{period_id}/export.xlsx"
    )
    workbook = load_workbook(BytesIO(response.content))
    row = list(workbook["Expenses"].iter_rows(values_only=True))[1]
    assert row[6:] == ("Rent", "Home", None, "Alex")


async def test_export_404_when_no_period(client):
    resp = await client.get("/api/v1/workspaces/1/periods/999/export.xlsx")
    assert resp.status_code == 404


async def test_export_income_type_and_running_balance(client):
    start = date.today()
    resp = await client.post(
        "/api/v1/workspaces/1/periods",
        json={
            "total_amount": "1000",
            "start_date": start.isoformat(),
            "end_date": (start + timedelta(days=9)).isoformat(),
        },
    )
    assert resp.status_code == 200
    period_id = resp.json()["period"]["id"]
    await client.post("/api/v1/workspaces/1/operations", json={"amount": "250"})
    await client.post("/api/v1/workspaces/1/operations", json={"amount": "100", "kind": "income"})

    resp = await client.get(f"/api/v1/workspaces/1/periods/{period_id}/export.xlsx")
    wb = load_workbook(BytesIO(resp.content))
    body = list(wb["Expenses"].iter_rows(values_only=True))[1:]
    assert [r[2] for r in body] == ["expense", "income"]
    # amounts stay positive in the sheet; the Type column carries the sign,
    # and the running balance applies it: 1000 - 250, then + 100
    assert [D(str(r[1])) for r in body] == [D("250"), D("100")]
    assert [D(str(r[4])) for r in body] == [D("750"), D("850")]

    kv = {row[0]: row[1] for row in wb["Period"].iter_rows(values_only=True)}
    assert D(str(kv["remaining"])) == D("850")


async def test_export_can_target_historical_period(client):
    today = date.today()
    old = {
        "total_amount": "300",
        "start_date": (today - timedelta(days=10)).isoformat(),
        "end_date": (today - timedelta(days=1)).isoformat(),
    }
    old_id = (await client.post("/api/v1/workspaces/1/periods", json=old)).json()["period"]["id"]
    await client.post(
        f"/api/v1/workspaces/1/periods/{old_id}/operations",
        json={"amount": "25", "occurred_on": old["end_date"]},
    )
    await client.post(
        "/api/v1/workspaces/1/periods",
        json={
            "total_amount": "1000",
            "start_date": today.isoformat(),
            "end_date": (today + timedelta(days=9)).isoformat(),
        },
    )

    response = await client.get(f"/api/v1/workspaces/1/periods/{old_id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    period_values = {
        row[0]: row[1] for row in workbook["Period"].iter_rows(values_only=True)
    }
    assert D(str(period_values["total_amount"])) == D("300")
    assert D(str(period_values["spent_total"])) == D("25")
